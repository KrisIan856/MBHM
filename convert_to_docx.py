#!/usr/bin/env python3
"""Convert Markdown documentation files to .docx using python-docx."""

import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)


def add_inline_runs(p, text):
    pattern = re.compile(r"(\*\*.+?\*\*|\*.+?\*)", re.DOTALL)
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run = p.add_run(part[1:-1])
            run.italic = True
        else:
            p.add_run(part)


def convert_markdown_to_docx(input_path, output_path):
    with open(input_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].rstrip("\n").strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            add_code_block(doc, "\n".join(code_lines))
            continue

        if stripped in ("---", "***", "___") and set(stripped) <= set("-*_"):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "000000")
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        m = re.match(r"^#\s+(.+)", stripped)
        if m:
            doc.add_heading(m.group(1), level=0)
            i += 1
            continue

        m = re.match(r"^##\s+(.+)", stripped)
        if m:
            doc.add_heading(m.group(1), level=1)
            i += 1
            continue

        m = re.match(r"^###\s+(.+)", stripped)
        if m:
            doc.add_heading(m.group(1), level=2)
            i += 1
            continue

        m = re.match(r"^####\s+(.+)", stripped)
        if m:
            doc.add_heading(m.group(1), level=3)
            i += 1
            continue

        if stripped.startswith("|"):
            table_rows = []
            while i < len(lines) and lines[i].rstrip("\n").strip().startswith("|"):
                row_line = lines[i].rstrip("\n").strip()
                if re.match(r"^\|[\s\-:|]+\|$", row_line):
                    i += 1
                    continue
                cells = [c.strip() for c in row_line.strip("|").split("|")]
                table_rows.append(cells)
                i += 1
            if table_rows:
                tbl = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                tbl.style = "Table Grid"
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row_cells in enumerate(table_rows):
                    for c_idx, cell_text in enumerate(row_cells):
                        cell = tbl.rows[r_idx].cells[c_idx]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        add_inline_runs(p, cell_text)
                        if r_idx == 0:
                            for run in p.runs:
                                run.bold = True
            continue

        m = re.match(r"^[-*+]\s+(.+)", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^\d+\.\s+(.+)", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, m.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, stripped)
        i += 1

    doc.save(output_path)
    print(f"Created {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python convert_to_docx.py <input.md> <output.docx>")
        sys.exit(1)
    convert_markdown_to_docx(sys.argv[1], sys.argv[2])
